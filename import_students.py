import sqlite3
from openpyxl import load_workbook


# ==========================================
# Configuration
# ==========================================

EXCEL_FILE = "students.xlsx"
DATABASE_FILE = "attendance.db"


# ==========================================
# Connect to database
# ==========================================

connection = sqlite3.connect(DATABASE_FILE)

cursor = connection.cursor()


# ==========================================
# Open Excel
# ==========================================

print()
print("==========================================")
print("       STUDENT EXCEL IMPORT")
print("==========================================")
print()

print(f"Opening: {EXCEL_FILE}")

workbook = load_workbook(
    EXCEL_FILE,
    read_only=True,
    data_only=True
)

sheet = workbook.active


# ==========================================
# Read Excel headers
# ==========================================

headers = []

for cell in sheet[1]:

    if cell.value is None:
        headers.append("")

    else:
        headers.append(
            str(cell.value).strip()
        )


print()
print("Excel columns:")

for header in headers:
    print("-", header)


# ==========================================
# Find column index
# ==========================================

def find_column(column_name):

    target = column_name.lower().strip()

    for index, header in enumerate(headers):

        if header.lower().strip() == target:

            return index

    return None


roll_col = find_column("Hall Ticket Number")
name_col = find_column("Name of the Student")
branch_col = find_column("Branch")
section_col = find_column("Section")
mobile_col = find_column("Mobile Number")
email_col = find_column("Emailid")


# ==========================================
# Check columns
# ==========================================

required = {
    "Hall Ticket Number": roll_col,
    "Name of the Student": name_col,
    "Branch": branch_col,
    "Section": section_col,
    "Mobile Number": mobile_col,
    "Emailid": email_col
}


missing = []

for name, index in required.items():

    if index is None:
        missing.append(name)


if missing:

    print()
    print("ERROR: Missing columns:")

    for column in missing:
        print("-", column)

    workbook.close()
    connection.close()
    exit()


# ==========================================
# Import counters
# ==========================================

new_students = 0
updated_students = 0
skipped_students = 0


# ==========================================
# Read student rows
# ==========================================

for row_number, row in enumerate(
    sheet.iter_rows(
        min_row=2,
        values_only=True
    ),
    start=2
):

    roll_number = row[roll_col]
    name = row[name_col]
    branch = row[branch_col]
    section = row[section_col]
    mobile = row[mobile_col]
    email = row[email_col]


    # --------------------------------------
    # Skip completely empty rows
    # --------------------------------------

    if roll_number is None and name is None:

        continue


    # --------------------------------------
    # Convert to strings
    # --------------------------------------

    if roll_number is not None:
        roll_number = str(roll_number).strip()

    if name is not None:
        name = str(name).strip()

    if branch is not None:
        branch = str(branch).strip()

    if section is not None:
        section = str(section).strip()

    if mobile is not None:
        mobile = str(mobile).strip()

    if email is not None:
        email = str(email).strip()


    # --------------------------------------
    # Validate roll number and name
    # --------------------------------------

    if not roll_number or not name:

        print(
            f"[SKIPPED] Excel row {row_number}: "
            "missing roll number or name"
        )

        skipped_students += 1

        continue


    # --------------------------------------
    # Default values
    # --------------------------------------

    if not branch:
        branch = "ECE"

    if not section:
        section = "A"


    # ======================================
    # Check if student already exists
    # ======================================

    cursor.execute(
        """
        SELECT id
        FROM students
        WHERE student_id = ?
        """,
        (roll_number,)
    )

    existing = cursor.fetchone()


    # ======================================
    # Update existing student
    # ======================================

    if existing:

        cursor.execute(
            """
            UPDATE students
            SET
                name = ?,
                email = ?,
                parent_mobile = ?,
                department = ?,
                year = ?,
                section = ?
            WHERE student_id = ?
            """,
            (
                name,
                email,
                mobile,
                branch,
                4,
                section,
                roll_number
            )
        )

        updated_students += 1

        print(
            f"[UPDATED] {roll_number} - {name}"
        )


    # ======================================
    # Insert new student
    # ======================================

    else:

        cursor.execute(
            """
            INSERT INTO students (
                student_id,
                name,
                email,
                parent_mobile,
                department,
                year,
                section
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                roll_number,
                name,
                email,
                mobile,
                branch,
                4,
                section
            )
        )

        new_students += 1

        print(
            f"[IMPORTED] {roll_number} - {name}"
        )


# ==========================================
# Save changes
# ==========================================

connection.commit()


# ==========================================
# Count students
# ==========================================

cursor.execute(
    "SELECT COUNT(*) FROM students"
)

total_students = cursor.fetchone()[0]


# ==========================================
# Close
# ==========================================

workbook.close()
connection.close()


# ==========================================
# Final result
# ==========================================

print()
print("==========================================")
print("          IMPORT COMPLETE")
print("==========================================")
print()

print(
    f"New students imported : {new_students}"
)

print(
    f"Existing students updated : {updated_students}"
)

print(
    f"Rows skipped : {skipped_students}"
)

print(
    f"Total students in database : {total_students}"
)

print()